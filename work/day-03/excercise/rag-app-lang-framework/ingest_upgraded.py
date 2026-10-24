from pathlib import Path

import chromadb
from docx import Document as WordDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

BASE_DIR = Path(__file__).resolve().parent
DOC_PATH = BASE_DIR / "medical_knowledge_dataset.docx"
CHROMA_PATH = BASE_DIR / "chroma_db"
COLLECTION_NAME = "medical_docs"

MIN_CHAR_LENGTH = 50
CHUNK_SIZE = 500
CHUNK_OVERLAP = 100
BATCH_SIZE = 500
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"


def load_chunks():
    if not DOC_PATH.exists():
        raise FileNotFoundError(f"Dataset not found: {DOC_PATH}")

    word_doc = WordDocument(str(DOC_PATH))
    print(f"Total paragraphs in document: {len(word_doc.paragraphs)}")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    chunks = []
    for paragraph_index, paragraph in enumerate(word_doc.paragraphs):
        text = paragraph.text.strip()
        if len(text) <= MIN_CHAR_LENGTH:
            continue
        source_doc = Document(
            page_content=text,
            metadata={"source": "medical_doc", "paragraph_index": paragraph_index},
        )
        chunks.extend(splitter.split_documents([source_doc]))

    for chunk_id, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = chunk_id

    print(f"Total chunks created: {len(chunks)}")
    return chunks


def main():
    documents = load_chunks()
    if not documents:
        print("No usable text found in the dataset; nothing to ingest.")
        return

    client = chromadb.PersistentClient(path=str(CHROMA_PATH))
    collection = client.get_or_create_collection(name=COLLECTION_NAME)
    if collection.count() > 0:
        print("Collection already contains data; skipping ingestion.")
        print(f"Existing document count: {collection.count()}")
        return

    vector_store = Chroma(
        client=client,
        collection_name=COLLECTION_NAME,
        embedding_function=HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL),
    )

    for start in range(0, len(documents), BATCH_SIZE):
        batch = documents[start : start + BATCH_SIZE]
        ids = [f"doc_{i}" for i in range(start, start + len(batch))]
        vector_store.add_documents(documents=batch, ids=ids)
        print(f"Ingested batch {start // BATCH_SIZE + 1}: {start} → {start + len(batch)}")

    print(f"\nIngestion complete! Total documents stored: {collection.count()}")


if __name__ == "__main__":
    main()